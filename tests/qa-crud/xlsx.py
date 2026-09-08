"""讀取／填寫 .xlsx 測試檔，供 tests/qa-crud 的匯入匯出驗證使用。

用法：
  python tests/qa-crud/xlsx.py dump <file.xlsx> [max_rows]
  python tests/qa-crud/xlsx.py fill <src.xlsx> <dst.xlsx> <rows.json> [start_row] [sheet]

rows.json 是二維陣列，每個內層陣列是一列的儲存格值（null 代表留白不寫）。
"""
import json
import sys

import openpyxl


def dump(path, max_rows=25):
    wb = openpyxl.load_workbook(path, data_only=False)
    out = {"sheets": []}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), values_only=True):
            rows.append(["" if c is None else str(c) for c in row])
        out["sheets"].append({
            "name": ws.title,
            "dimensions": ws.dimensions,
            "maxRow": ws.max_row,
            "maxCol": ws.max_column,
            "rows": rows,
        })
    print(json.dumps(out, ensure_ascii=False, indent=1))


def fill(src, dst, rows_json, start_row, sheet):
    wb = openpyxl.load_workbook(src)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = json.loads(rows_json) if rows_json.strip().startswith("[") else json.load(open(rows_json, encoding="utf-8"))
    for r, values in enumerate(rows, start=start_row):
        for c, value in enumerate(values, start=1):
            if value is None:
                continue
            ws.cell(row=r, column=c, value=value)
    wb.save(dst)
    print(json.dumps({"saved": dst, "rows": len(rows), "startRow": start_row, "sheet": ws.title}, ensure_ascii=False))


def create(dst, rows_json, sheet_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    rows = json.loads(rows_json) if rows_json.strip().startswith("[") else json.load(open(rows_json, encoding="utf-8"))
    for values in rows:
        ws.append(values)
    wb.save(dst)
    print(json.dumps({"saved": dst, "rows": len(rows), "sheet": ws.title}, ensure_ascii=False))


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "dump":
        dump(sys.argv[2], int(sys.argv[3]) if len(sys.argv) > 3 else 25)
    elif cmd == "create":
        create(sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else "Sheet1")
    elif cmd == "fill":
        fill(sys.argv[2], sys.argv[3], sys.argv[4],
             int(sys.argv[5]) if len(sys.argv) > 5 else 2,
             sys.argv[6] if len(sys.argv) > 6 else None)
    else:
        raise SystemExit(__doc__)
